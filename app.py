"""
app.py — Streamlit Application
AI Meeting Minutes Summarizer
Auto-trains the model on first launch if not already trained.
"""

import os
import sys
import io

import numpy as np
import streamlit as st

# ─── Path setup ───────────────────────────────────────────────────────────────
BASE_DIR       = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)

MODEL_PATH     = os.path.join(BASE_DIR, "models", "seq2seq_model.h5")
TOKENIZER_PATH = os.path.join(BASE_DIR, "models", "tokenizers.pkl")
DATA_DIR       = os.path.join(BASE_DIR, "Data")

# ─── Page config ──────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="AI Meeting Summarizer",
    page_icon="🤖",
    layout="centered",
)


# ═════════════════════════════════════════════════════════════════════════════
# AUTO-TRAIN on first launch
# ═════════════════════════════════════════════════════════════════════════════

def run_training():
    """Full training pipeline — runs once on Streamlit Cloud if model is missing."""
    from src.eda           import load_data, run_eda
    from src.preprocessing import (preprocess, MAX_TEXT_LEN, MAX_SUMMARY_LEN,
                                    VOCAB_SIZE_TEXT, VOCAB_SIZE_SUM)
    from src.model         import build_seq2seq, train_model

    os.makedirs(os.path.join(BASE_DIR, "models"), exist_ok=True)
    os.makedirs(os.path.join(BASE_DIR, "plots"),  exist_ok=True)

    st.info("⏳ Step 1/4 — Loading & analysing data…")
    df = load_data(DATA_DIR)
    _, df = run_eda(df)

    st.info("⏳ Step 2/4 — Preprocessing text…")
    x_train, x_val, y_train, y_val, text_tok, sum_tok = preprocess(
        df, tokenizer_path=TOKENIZER_PATH
    )

    text_vocab = min(VOCAB_SIZE_TEXT, len(text_tok.word_index) + 1)
    sum_vocab  = min(VOCAB_SIZE_SUM,  len(sum_tok.word_index)  + 1)

    st.info("⏳ Step 3/4 — Building Seq2Seq model…")
    model = build_seq2seq(
        text_vocab=text_vocab,
        sum_vocab=sum_vocab,
        max_text_len=MAX_TEXT_LEN,
        max_summary_len=MAX_SUMMARY_LEN,
    )

    st.info("⏳ Step 4/4 — Training model (this takes a few minutes)…")
    train_model(
        model, x_train, y_train, x_val, y_val,
        batch_size=64, epochs=30,
        model_path=MODEL_PATH,
    )
    st.success("✅ Training complete! Reloading app…")
    st.rerun()


# Run training automatically if model doesn't exist yet
if not os.path.exists(MODEL_PATH):
    st.markdown(
        "<h1 style='text-align:center;color:#4F8BF9;'>🤖 AI Meeting Minutes Summarizer</h1>",
        unsafe_allow_html=True,
    )
    st.warning("🔧 **First launch detected.** Training the model now — please wait…")
    run_training()
    st.stop()


# ═════════════════════════════════════════════════════════════════════════════
# LOAD MODEL (cached after training)
# ═════════════════════════════════════════════════════════════════════════════

@st.cache_resource(show_spinner="Loading AI model…")
def load_model_cached():
    from tensorflow.keras.models import load_model as keras_load
    from src.model         import build_inference_models
    from src.preprocessing import load_tokenizers

    seq2seq = keras_load(MODEL_PATH)
    enc_model, dec_model = build_inference_models(seq2seq)
    text_tok, sum_tok    = load_tokenizers(TOKENIZER_PATH)
    return enc_model, dec_model, text_tok, sum_tok


# ═════════════════════════════════════════════════════════════════════════════
# HELPER FUNCTIONS
# ═════════════════════════════════════════════════════════════════════════════

def do_summarize(text: str) -> str:
    from src.summarizer import generate_summary
    enc_model, dec_model, text_tok, sum_tok = load_model_cached()
    summary = generate_summary(text, enc_model, dec_model, text_tok, sum_tok)
    return summary if summary.strip() else "Unable to generate summary. Try a longer input."


def compute_stats(original: str, summary: str) -> dict:
    orig  = len(original.split())
    summ  = len(summary.split())
    ratio = round((1 - summ / orig) * 100, 1) if orig > 0 else 0
    return {
        "Original Words":    orig,
        "Summary Words":     summ,
        "Compression Ratio": f"{ratio}%",
    }


def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        import fitz
        doc  = fitz.open(stream=file_bytes, filetype="pdf")
        return "\n".join(page.get_text() for page in doc).strip()
    except Exception as e:
        return f"[Error reading PDF: {e}]"


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
        doc  = Document(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip()).strip()
    except Exception as e:
        return f"[Error reading DOCX: {e}]"


def create_pdf_download(summary: str) -> bytes:
    try:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, "AI Generated Meeting Summary", ln=True, align="C")
        pdf.ln(4)
        pdf.set_draw_color(100, 100, 100)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
        pdf.ln(6)
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 8, summary)
        return bytes(pdf.output())
    except Exception as e:
        return f"[PDF error: {e}]".encode()


# ═════════════════════════════════════════════════════════════════════════════
# SESSION STATE
# ═════════════════════════════════════════════════════════════════════════════

if "summary"     not in st.session_state: st.session_state["summary"]     = ""
if "source_text" not in st.session_state: st.session_state["source_text"] = ""


# ═════════════════════════════════════════════════════════════════════════════
# HEADER
# ═════════════════════════════════════════════════════════════════════════════

st.markdown(
    """
    <h1 style='text-align:center;color:#4F8BF9;'>🤖 AI Meeting Minutes Summarizer</h1>
    <p style='text-align:center;color:gray;'>Seq2Seq Encoder-Decoder · LSTM · Streamlit</p>
    <hr/>
    """,
    unsafe_allow_html=True,
)

st.success("✅ Model loaded and ready.", icon="✅")
st.markdown("---")


# ═════════════════════════════════════════════════════════════════════════════
# SECTION 1 — Enter Meeting Notes
# ═════════════════════════════════════════════════════════════════════════════

st.markdown("## 📝 Section 1 · Enter Meeting Notes")

meeting_text = st.text_area(
    label="Paste your meeting notes here:",
    placeholder=(
        "e.g.\nYesterday the development team completed the authentication APIs. "
        "The frontend team integrated the login functionality. "
        "Testing will begin next Monday..."
    ),
    height=220,
    key="manual_input",
)

st.markdown("---")


# ═════════════════════════════════════════════════════════════════════════════
# SECTION 2 — Generate Summary Button
# ═════════════════════════════════════════════════════════════════════════════

st.markdown("## ▶️ Section 2 · Generate Summary")

if st.button("🚀 Generate Summary", use_container_width=True, type="primary"):
    if meeting_text.strip():
        with st.spinner("Generating summary…"):
            st.session_state["summary"]     = do_summarize(meeting_text)
            st.session_state["source_text"] = meeting_text
    else:
        st.warning("Please paste some meeting notes in Section 1 first.")

st.markdown("---")


# ═════════════════════════════════════════════════════════════════════════════
# SECTION 3 — AI Generated Summary
# ═════════════════════════════════════════════════════════════════════════════

st.markdown("## 📋 Section 3 · AI Generated Summary")

if st.session_state["summary"]:
    st.success(st.session_state["summary"])
else:
    st.info("Your AI-generated summary will appear here after clicking **Generate Summary**.")

st.markdown("---")


# ═════════════════════════════════════════════════════════════════════════════
# SECTION 4 — Statistics
# ═════════════════════════════════════════════════════════════════════════════

st.markdown("## 📊 Section 4 · Statistics")

if st.session_state["summary"] and st.session_state["source_text"]:
    stats = compute_stats(st.session_state["source_text"], st.session_state["summary"])
    c1, c2, c3 = st.columns(3)
    c1.metric("Original Words",    stats["Original Words"])
    c2.metric("Summary Words",     stats["Summary Words"])
    c3.metric("Compression Ratio", stats["Compression Ratio"])
else:
    st.info("Statistics will appear here once a summary is generated.")

st.markdown("---")


# ═════════════════════════════════════════════════════════════════════════════
# BONUS FEATURES
# ═════════════════════════════════════════════════════════════════════════════

st.markdown("## 🎁 Bonus Features")

tab1, tab2, tab3 = st.tabs(["📄 Upload PDF", "📝 Upload DOCX", "⬇️ Download Summary"])

# ── Bonus 1: PDF ──────────────────────────────────────────────────────────────
with tab1:
    st.markdown("### Upload a PDF Meeting Document")
    pdf_file = st.file_uploader("Choose a PDF file", type=["pdf"], key="pdf_uploader")
    if pdf_file:
        pdf_text = extract_text_from_pdf(pdf_file.read())
        if pdf_text.startswith("[Error"):
            st.error(pdf_text)
        else:
            st.text_area("Extracted Text (preview):",
                         value=pdf_text[:1000] + "…" if len(pdf_text) > 1000 else pdf_text,
                         height=160, disabled=True)
            if st.button("🚀 Summarize PDF", key="pdf_btn", type="primary"):
                with st.spinner("Summarizing…"):
                    s = do_summarize(pdf_text)
                    st.session_state["summary"]     = s
                    st.session_state["source_text"] = pdf_text
                st.success(s)
                stats = compute_stats(pdf_text, s)
                c1, c2, c3 = st.columns(3)
                c1.metric("Original Words",    stats["Original Words"])
                c2.metric("Summary Words",     stats["Summary Words"])
                c3.metric("Compression Ratio", stats["Compression Ratio"])

# ── Bonus 2: DOCX ─────────────────────────────────────────────────────────────
with tab2:
    st.markdown("### Upload a DOCX Meeting Document")
    docx_file = st.file_uploader("Choose a DOCX file", type=["docx"], key="docx_uploader")
    if docx_file:
        docx_text = extract_text_from_docx(docx_file.read())
        if docx_text.startswith("[Error"):
            st.error(docx_text)
        else:
            st.text_area("Extracted Text (preview):",
                         value=docx_text[:1000] + "…" if len(docx_text) > 1000 else docx_text,
                         height=160, disabled=True)
            if st.button("🚀 Summarize DOCX", key="docx_btn", type="primary"):
                with st.spinner("Summarizing…"):
                    s = do_summarize(docx_text)
                    st.session_state["summary"]     = s
                    st.session_state["source_text"] = docx_text
                st.success(s)
                stats = compute_stats(docx_text, s)
                c1, c2, c3 = st.columns(3)
                c1.metric("Original Words",    stats["Original Words"])
                c2.metric("Summary Words",     stats["Summary Words"])
                c3.metric("Compression Ratio", stats["Compression Ratio"])

# ── Bonus 3: Download as PDF ──────────────────────────────────────────────────
with tab3:
    st.markdown("### Download Your Summary as PDF")
    if st.session_state["summary"]:
        pdf_bytes = create_pdf_download(st.session_state["summary"])
        st.download_button(
            label="⬇️ Download Summary PDF",
            data=pdf_bytes,
            file_name="meeting_summary.pdf",
            mime="application/pdf",
            use_container_width=True,
            type="primary",
        )
    else:
        st.info("Generate a summary first, then download it as PDF here.")


# ─── Footer ───────────────────────────────────────────────────────────────────
st.markdown("---")
st.markdown(
    "<p style='text-align:center;color:gray;font-size:13px;'>"
    "AI Meeting Minutes Summarizer · Seq2Seq LSTM · Built with Streamlit"
    "</p>",
    unsafe_allow_html=True,
)
